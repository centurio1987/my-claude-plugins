#!/usr/bin/env python3
"""플레이스홀더 린터 — 채우다 만 자리가 발행물에 남는 것을 막는다.

규약(`std-v1`)은 하나지만 **표기는 파일 타입별로 갈린다.** 억지로 한 표기로 밀면 깨진다:
  · `[대괄호]` 는 마크다운 링크 문법과 충돌한다.
  · `{{중괄호}}` 는 MDX 에서 JSX 표현식으로 파싱돼 빌드가 죽는다.
그래서 id 문법과 규약 이름은 하나로 두고 표기만 이원화한다.

| 용도      | .md / .html / .txt / .pptx        | .mdx                        |
|-----------|-----------------------------------|-----------------------------|
| 채움 슬롯 | `{{ id }}`                        | `{/* @slot:id */}`          |
| 집필 지시 | `<!-- @write: … -->`              | `{/* @write: … */}`         |
| 조건 블록 | `<!-- @if: k=v -->…<!-- @end -->` | `{/* @if: k=v */}…{/* @end */}` |
| 메타 힌트 | `<!-- @hint: … -->`               | `{/* @hint: … */}`          |
| 집필 마커 | `@[제안]` `@[추정]` `@[창작]` `@[what-if]` `@[확인]` | 동일 |

**파서는 포맷마다 다르다.** 정규식 하나로 넷을 다루면 반드시 깨진다 —
특히 PPTX 는 한 문장이 여러 텍스트 run 으로 쪼개져 저장돼서 `{{ id }}` 가 조각난다.

사용법:
    lint_placeholders.py check <파일> [--spec <id>] [--project <dir>] --stage draft|final
    lint_placeholders.py import <템플릿> [--out template.slots.json]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

# ── std-v1 정규식 ────────────────────────────────────────────────────────
SLOT_ID = r"[\w.:-]+"
RE_SLOT_BRACE = re.compile(r"\{\{\s*(" + SLOT_ID + r")\s*\}\}")
RE_SLOT_MDX = re.compile(r"\{\s*/\*\s*@slot:\s*(" + SLOT_ID + r")\s*\*/\s*\}")
RE_WRITE_HTML = re.compile(r"<!--\s*@write:(.*?)-->", re.S)
RE_WRITE_MDX = re.compile(r"\{\s*/\*\s*@write:(.*?)\*/\s*\}", re.S)
RE_IF_HTML = re.compile(r"<!--\s*@if:\s*([^>]*?)-->")
RE_IF_MDX = re.compile(r"\{\s*/\*\s*@if:\s*(.*?)\*/\s*\}")
RE_END_HTML = re.compile(r"<!--\s*@end\s*-->")
RE_END_MDX = re.compile(r"\{\s*/\*\s*@end\s*\*/\s*\}")
RE_HINT = re.compile(r"(?:<!--|\{\s*/\*)\s*@hint:")
RE_MARKER = re.compile(r"@\[(제안|추정|창작|what-if|확인)\]")
# std-v1 은 `@[추정]` 이지만 resume 는 이관 전부터 `[추정]` 을 써 왔고, 그 표기가
# `write-resume` 절차와 기존 draft 전체에 박혀 있다. 레거시 표기를 못 보면
# **마커가 있는 문서를 "마커 0" 이라고 보고**한다 — 거짓 실패보다 나쁜 거짓 통과다.
# D6 대로 레거시는 인식하되 통일 대상은 아니다: draft 에서는 경고, final 에서는 실패.
RE_MARKER_LEGACY = re.compile(r"(?<!@)\[(제안|추정|창작|what-if|확인)\]")

# 레거시 — 인식만 하고 경고. 기존 산출물을 건드리지 않기 위한 것이다.
LEGACY = {
    "bracket-inline": re.compile(r"(?<!\!)\[[^\]\n]{1,40}\](?!\()"),   # resume 슬롯. 링크는 제외
    "pos-block": re.compile(r"\{\{pos:[\w-]+\}\}|<!--\s*/?pos:"),      # resume 직무 분기
    "brace-descriptive": re.compile(r"\{[가-힣][^{}\n]{1,30}\}"),      # code_test 캔버스
    "ordinal": re.compile(r"\bs\d{4}\b"),                              # resume 템플릿 슬롯 키
}

# `[…]` 는 슬롯 표기이기도 하지만 **배열 리터럴·인덱스이기도 하다.**
# 알고리즘 글에서는 후자가 압도적이라(`[3, 2, 2, 4, 2]`, `[i]`, `[K]`) 걸러내지 않으면
# 정상 본문 수십 건이 미해결 슬롯으로 잡혀 발행이 막힌다. 실제로 그렇게 막혔다.
ARRAY_LITERAL = re.compile(r"^[\s\d,.\-+*/()…N+]*$")   # 숫자·연산자·공백만
SHORT_INDEX = re.compile(r"^[A-Za-z0-9]{1,2}$")        # [i] [K] [3] 같은 인덱스


def _looks_like_slot(inner: str) -> bool:
    """`[...]` 안의 내용이 **채워야 할 자리**로 보이는가.

    슬롯은 사람이 읽는 라벨이다(`[이름]`, `[한 줄 요약]`, `[s0001]`).
    배열 리터럴과 인덱스는 그 글의 내용이지 빈자리가 아니다.
    """
    s = inner.strip()
    if not s or SHORT_INDEX.match(s):
        return False
    if ARRAY_LITERAL.match(s) and any(ch.isdigit() for ch in s):
        return False
    return True

FENCE = re.compile(r"^\s*(```|~~~)")


# ── 파일 타입별 어댑터 ───────────────────────────────────────────────────
def text_of_md(path: Path) -> str:
    """코드펜스 안은 비운다 — 예시 코드에 든 중괄호는 슬롯이 아니다."""
    out, in_fence = [], False
    for line in path.read_text(encoding="utf-8").splitlines():
        if FENCE.match(line):
            in_fence = not in_fence
            out.append("")
            continue
        out.append("" if in_fence else line)
    return "\n".join(out)


def text_of_html(path: Path) -> str:
    """텍스트 노드만 모은다. 6MB 번들이 올 수 있으니 태그를 통째로 지운다."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"<(script|style)\b.*?</\1>", " ", raw, flags=re.S | re.I)
    # HTML 주석은 지우지 않는다 — @write·@if 가 주석으로 산다.
    return re.sub(r"<(?!!--)[^>]*>", " ", raw)


class UnreadableFile(Exception):
    """읽지 못한 파일을 '깨끗하다'고 말하지 않기 위한 예외.

    의존성이 없어서 빈 문자열을 돌려주면 린터가 '슬롯 0개, 통과'라고 보고한다.
    못 읽은 것과 깨끗한 것은 완전히 다른 결과다 — 조용한 통과가 가장 나쁜 실패다.
    """


def text_of_pptx(path: Path) -> str:
    """shape·run 을 순회하되 **run 을 문단 단위로 이어 붙인 뒤** 검출한다.

    PowerPoint 는 한 문장을 서식 단위로 쪼개 저장한다. `{{ header.title }}` 이
    `{{ head`/`er.ti`/`tle }}` 세 run 으로 나뉘는 일이 흔해서, run 마다 정규식을 돌리면
    아무것도 못 찾는다. 이어 붙인 뒤 봐야 한다.
    """
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        raise UnreadableFile(
            ".pptx 를 열려면 python-pptx 가 필요하다 (`pip install python-pptx`). "
            "설치 전에는 이 파일을 검사한 것으로 치지 않는다.")
    chunks = []
    prs = Presentation(str(path))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                chunks.append("".join(r.text for r in para.runs))
        if slide.has_notes_slide:
            chunks.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(chunks)


def read_text(path: Path) -> tuple[str, bool]:
    """(검사할 텍스트, mdx 인가)."""
    suf = path.suffix.lower()
    if suf == ".mdx":
        return text_of_md(path), True
    if suf in (".md", ".txt"):
        return text_of_md(path), False
    if suf in (".html", ".htm"):
        return text_of_html(path), False
    if suf == ".pptx":
        return text_of_pptx(path), False
    return path.read_text(encoding="utf-8", errors="replace"), False


# ── 검사 ─────────────────────────────────────────────────────────────────
def scan(text: str, is_mdx: bool) -> dict:
    slots = RE_SLOT_MDX.findall(text) if is_mdx else RE_SLOT_BRACE.findall(text)
    # .md 계열도 혹시 mdx 표기가 섞였으면 함께 잡는다(둘 다 미해결이라는 사실은 같다).
    slots += RE_SLOT_MDX.findall(text) if not is_mdx else RE_SLOT_BRACE.findall(text)
    writes = (RE_WRITE_MDX if is_mdx else RE_WRITE_HTML).findall(text)
    ifs = (RE_IF_MDX if is_mdx else RE_IF_HTML).findall(text)
    ends = len((RE_END_MDX if is_mdx else RE_END_HTML).findall(text))
    return {
        "slots": sorted(set(slots)),
        "writes": len(writes),
        "hints": len(RE_HINT.findall(text)),
        "markers": sorted(set(RE_MARKER.findall(text))),
        "markers_legacy": sorted(set(RE_MARKER_LEGACY.findall(text))),
        "if_open": len(ifs),
        "if_close": ends,
        "legacy": _legacy_counts(text),
    }


def _legacy_counts(text: str) -> dict:
    out = {}
    for k, pat in LEGACY.items():
        hits = pat.findall(text)
        if k == "bracket-inline":
            hits = [h for h in hits if _looks_like_slot(h.strip("[]"))]
        if hits:
            out[k] = len(hits)
    return out


def deferred_slots(slots: list[str], accepted: list[str]) -> tuple[list[str], list[str]]:
    """(다음 단계로 미룬 슬롯, 그냥 남은 슬롯).

    파이프라인이 draft → set → format 세 단계인 곳이 있다. resume 가 그렇다:
    `03_set` 은 직무 공통 1벌이라 `{{pos:title}}` 같은 **직무 분기 슬롯을 일부러 남기고**,
    4차 형식화가 직무별로 해소한다. 그 슬롯이 없으면 렌더 계약이 깨진다.

    그런데 set 을 `--stage final` 로 재면 바로 그 슬롯이 실패로 잡힌다. 명세는
    `placeholders.legacy_accept` 로 그 표기를 **명시적으로 허용**하고 있는데도 그렇다.
    산출물이 틀린 게 아니라 단계 모델이 두 칸뿐이었던 것이다.

    무엇을 미룰 수 있는지는 도구가 아니라 **명세가 정한다** — 선언 안 한 슬롯은 그냥 남은 것이다.
    """
    pats = [LEGACY[k] for k in accepted if k in LEGACY]
    deferred, stray = [], []
    for sid in slots:
        rendered = "{{" + sid + "}}"
        (deferred if any(pat.search(rendered) for pat in pats) else stray).append(sid)
    return deferred, stray


def cmd_check(args) -> int:
    path = Path(args.path).expanduser()
    if not path.exists():
        print(f"파일이 없다: {path}", file=sys.stderr)
        return 2
    try:
        text, is_mdx = read_text(path)
    except UnreadableFile as e:
        print(f"# 검사 불가: {path.name}", file=sys.stderr)
        print(f"  {e}", file=sys.stderr)
        return 2                      # 통과(0)도 위반(1)도 아니다 — 아예 재지 못했다
    r = scan(text, is_mdx)

    accepted: list[str] = []
    if args.spec:
        sp = Path(args.project or ".").expanduser() / ".claude/authoring/specs" / args.spec / "spec.json"
        if sp.exists():
            accepted = json.loads(sp.read_text(encoding="utf-8")).get(
                "placeholders", {}).get("legacy_accept", [])

    print(f"# 플레이스홀더 검사: {path.name}  (stage={args.stage}, {'mdx' if is_mdx else path.suffix or 'text'})")
    print(f"  미해결 슬롯   {len(r['slots'])}  {r['slots'][:8]}")
    print(f"  집필 지시     {r['writes']}")
    print(f"  메타 힌트     {r['hints']}")
    print(f"  집필 마커     {len(r['markers'])}  {r['markers']}")
    if r.get("markers_legacy"):
        print(f"  └ 레거시 표기 {len(r['markers_legacy'])}  {r['markers_legacy']}"
              "  (std-v1 은 `@[…]` — 기존 산출물이라 통일 대상은 아니다)")
    print(f"  조건 블록     열림 {r['if_open']} / 닫힘 {r['if_close']}")
    if r["legacy"]:
        for k, n in r["legacy"].items():
            tag = "허용(spec)" if k in accepted else "경고"
            print(f"  레거시 {k:20s} {n:4d}건  [{tag}]")

    fail = []
    if r["if_open"] != r["if_close"]:
        fail.append(f"조건 블록 짝이 안 맞는다 (열림 {r['if_open']} / 닫힘 {r['if_close']})")
    if args.stage in ("set", "final"):
        # 완성본에는 아무것도 남으면 안 된다. draft 단계에서는 지시·마커가 정상이다.
        # set 은 그 중간이다 — 지시·마커는 이미 0 이어야 하고, 슬롯만 다음 단계로 미룰 수 있다.
        if args.stage == "final" and r["slots"]:
            fail.append(f"미해결 슬롯 {len(r['slots'])}개: {r['slots'][:5]}")
        elif args.stage == "set" and r["slots"]:
            deferred, stray = deferred_slots(r["slots"], accepted)
            if deferred:
                print(f"  └ 다음 단계로 미룸 {len(deferred)}  {deferred}  (spec 의 legacy_accept)")
            if stray:
                fail.append(f"미해결 슬롯 {len(stray)}개: {stray[:5]}")
        if r["writes"]:
            fail.append(f"집필 지시 {r['writes']}개가 남았다")
        if r["hints"]:
            fail.append(f"메타 힌트 {r['hints']}개가 남았다")
        if r["markers"]:
            fail.append(f"집필 마커가 남았다: {r['markers']}")
        # final 에서는 표기와 무관하게 0 이어야 한다. 레거시라고 봐주면 마커가 발행된다.
        if r.get("markers_legacy"):
            fail.append(f"집필 마커(레거시 표기)가 남았다: {r['markers_legacy']}")
        # 레거시는 **경고까지만** 한다. 정책이 "신규만 통일, 기존 산출물은 건드리지 않는다"이므로
        # 미선언 레거시를 실패로 처리하면 그 정책과 모순된다. 게다가 레거시 패턴은 넓어서
        # 오탐이 섞인다 — 실패로 만들면 정상 문서가 발행을 못 하게 된다(실제로 그랬다).
        undeclared = [k for k in r["legacy"] if k not in accepted]
        if undeclared:
            print(f"  경고: spec 에 선언되지 않은 레거시 표기 {undeclared} — "
                  f"의도한 것이면 placeholders.legacy_accept 에 넣어 두라")

    print()
    if fail:
        for f in fail:
            print(f"  실패: {f}")
        return 1
    print("  통과.")
    return 0


def cmd_import(args) -> int:
    """템플릿에서 슬롯을 뽑아 template.slots.json 을 만든다."""
    path = Path(args.path).expanduser()
    try:
        text, is_mdx = read_text(path)
    except UnreadableFile as e:
        print(f"# 수입 불가: {path.name}\n  {e}", file=sys.stderr)
        return 2
    r = scan(text, is_mdx)
    legacy_kind = max(r["legacy"], key=lambda k: r["legacy"][k]) if r["legacy"] else None
    out = {
        "schema_version": 1,
        "template": str(path),
        "format": "mdx" if is_mdx else (path.suffix.lstrip(".") or "text"),
        "slot_style": "std-v1" if r["slots"] else (legacy_kind or "none"),
        "slots": [{"id": s, "kind": "fill", "value": None} for s in r["slots"]],
        "legacy_detected": r["legacy"],
        "_note": "value 가 null 이면 미해결이다. 템플릿 파일은 프로젝트에 그대로 두고 paths.json 이 경로만 참조한다.",
    }
    dest = Path(args.out).expanduser() if args.out else path.with_name("template.slots.json")
    dest.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"슬롯 {len(r['slots'])}개 → {dest}")
    if r["legacy"]:
        print(f"  레거시 표기 감지: {r['legacy']} — spec 의 placeholders.legacy_accept 에 넣어 두면 통과한다.")
    if not r["slots"] and not r["legacy"]:
        print("  슬롯을 찾지 못했다. 템플릿이 std-v1 표기를 쓰는지 확인하라.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    c = sub.add_parser("check")
    c.add_argument("path")
    c.add_argument("--spec")
    c.add_argument("--project")
    c.add_argument("--stage", choices=["draft", "set", "final"], default="draft")
    c.set_defaults(func=cmd_check)

    i = sub.add_parser("import")
    i.add_argument("path")
    i.add_argument("--out")
    i.set_defaults(func=cmd_import)

    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
